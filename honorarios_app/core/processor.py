from openpyxl import Workbook

from honorarios_app.core.common import find_header_row, get_col, norm
from honorarios_app.core.excel_utils import (
    build_hosvireport,
    build_pct_override_if_present,
    build_tarifario_coosalud,
    load_excel,
    validate_required_sheets,
)
from honorarios_app.core.models import ProcessResult
from honorarios_app.core.output_writer import (
    add_pedir_fact_summary_box,
    add_pre_summary_box,
    autofit_columns,
    style_header,
)
from honorarios_app.services.specialty_factory import get_specialty


def process_excel(
    input_path: str,
    output_path: str,
    especialidad="ortopedia",
    original_sheet=None,
    hosvi_sheet="HOSVIREPORT",
    coosalud_sheet="COOSALUD",
    progress_cb=None,
    payment_pct=None,
    payment_cfg=None,
) -> ProcessResult:
    def emit(msg: str, frac: float):
        if progress_cb:
            try:
                progress_cb(msg, max(0.0, min(1.0, frac)))
            except Exception:
                pass

    specialty = get_specialty(especialidad)
    requires_hosvireport = getattr(
        specialty,
        "requires_hosvireport",
        especialidad != "cardiologia",
    )

    # Configuración de porcentajes
    if payment_cfg is None:
        if especialidad == "cardiologia":
            payment_cfg = {
                "consulta_pct": getattr(specialty, "default_consulta_pct", 0.90),
                "procedimiento_pct": getattr(specialty, "default_procedimiento_pct", 0.36),
            }
        else:
            if payment_pct is None:
                payment_pct = getattr(specialty, "default_payment_pct", 0.70)
            payment_cfg = {
                "payment_pct": payment_pct
            }

    emit("Abriendo Excel...", 0.02)
    wb = load_excel(input_path)
    ws_original = wb[original_sheet] if original_sheet else wb[wb.sheetnames[0]]

    # Hojas obligatorias
    if requires_hosvireport:
        if hosvi_sheet not in wb.sheetnames:
            raise RuntimeError(f"No existe la hoja '{hosvi_sheet}' en el archivo.")

    if coosalud_sheet not in wb.sheetnames:
        raise RuntimeError(f"No existe la hoja '{coosalud_sheet}' en el archivo.")

    emit("Leyendo tarifas COOSALUD...", 0.08)
    pct_override = build_pct_override_if_present(wb, "PORCENTAJE_PAGO")
    tarifario = build_tarifario_coosalud(wb[coosalud_sheet])

    hosvi_by_fact = {}
    hosvi_by_alt = {}

    if requires_hosvireport:
        emit("Leyendo HOSVIREPORT...", 0.14)
        hosvi_by_fact, hosvi_by_alt = build_hosvireport(wb[hosvi_sheet], pct_override=pct_override)

    hosvi_cursor_fact = {}
    hosvi_cursor_alt = {}

    emit("Detectando encabezados hoja original...", 0.18)
    header_row, col_map = find_header_row(ws_original, ["CONTR", "FAC", "PROCED", "COD"])
    if header_row is None:
        raise RuntimeError(
            "No pude detectar encabezados en hoja original "
            "(CONTRATO / NO FAC / PROCEDIMIENTO / CODIGO PROCEDIMIENTO)."
        )

    col_no_fac = get_col(col_map, "NO FAC", "NUM FAC", "N° FAC", "FACTURA")
    col_contrato = get_col(col_map, "CONTRATO")
    col_cod_proc = get_col(
        col_map,
        "CODIGO PROCEDIMIENTO",
        "COD PROCEDIMIENTO",
        "CODIGO PROC",
        "COD_PROCED",
        "COD",
    )
    col_proc = get_col(col_map, "PROCEDIMIENTO")
    col_saldo = get_col(col_map, "SALDO")
    col_estado = get_col(col_map, "ESTADO")
    col_tipo_fac = get_col(col_map, "TIPO FAC", "TIPO_FAC", "TIPO FACTURA")
    col_id_paciente = get_col(
        col_map,
        "IDENTIFICACION PACIENTE",
        "IDENTIFICACION_PACIENTE",
        "ID PACIENTE",
        "DOC PACIENTE",
        "DOCUMENTO PACIENTE",
        "IDENTIFICACION",
    )
    col_nombre_paciente = get_col(
        col_map,
        "NOMBRE PACIENTE",
        "NOMBRE_PACIENTE",
        "NOMBRE P.",
        "PACIENTE",
    )

    if any(c is None for c in [col_no_fac, col_contrato, col_cod_proc]):
        raise RuntimeError(
            "Faltan columnas obligatorias en hoja original: "
            "NO FAC, CONTRATO, CODIGO PROCEDIMIENTO."
        )

    base_headers = [ws_original.cell(header_row, c).value for c in range(1, ws_original.max_column + 1)]
    base_header_norm = [norm(h) for h in base_headers]

    idx_saldo = next((i for i, h in enumerate(base_header_norm, start=1) if h == "SALDO"), None)

    idx_vlr_aut = None
    try:
        idx_vlr_aut = base_header_norm.index("VLR. A AUTORIZAR") + 1
    except ValueError:
        for i, h in enumerate(base_header_norm, start=1):
            if "VLR" in h and "AUTORIZAR" in h:
                idx_vlr_aut = i
                break

    idx_total_coosalud = None

    emit("Creando archivo de salida...", 0.22)
    out = Workbook()
    ws_pre = out.active
    ws_pre.title = "PRE"
    ws_amar = out.create_sheet("AMARILLO")
    ws_pedir = out.create_sheet("PEDIR FAC")
    ws_anul = out.create_sheet("ANULADOS")

    extra_sheets = {}
    if hasattr(specialty, "get_extra_sheets"):
        extra_sheet_defs = specialty.get_extra_sheets() or {}
        for sheet_name, reason_header in extra_sheet_defs.items():
            ws_extra = out.create_sheet(sheet_name)
            ws_extra.append(base_headers + [reason_header])
            style_header(ws_extra)
            extra_sheets[sheet_name] = ws_extra

    # Encabezados calculados por especialidad
    if especialidad == "cardiologia":
        calc_headers = specialty.get_calc_headers(payment_cfg=payment_cfg)
    else:
        calc_headers = specialty.get_calc_headers(payment_pct=payment_cfg.get("payment_pct"))

    try:
        idx_total_coosalud = len(base_headers) + calc_headers.index("TOTAL COOSALUD") + 1
    except ValueError:
        idx_total_coosalud = None

    ws_pre.append(base_headers + calc_headers)
    ws_amar.append(base_headers + ["RAZON EXCLUSION"])
    ws_pedir.append(base_headers)
    ws_anul.append(base_headers + ["RAZON ANULADO"])

    for ws in (ws_pre, ws_amar, ws_pedir, ws_anul):
        style_header(ws)

    emit("Detectando anulaciones...", 0.24)
    anulable_rows = set()
    if hasattr(specialty, "detect_anulable_rows"):
        anulable_rows = specialty.detect_anulable_rows(ws_original, header_row, col_map)

    duplicate_rows = {}
    if hasattr(specialty, "detect_duplicate_rows"):
        duplicate_rows = specialty.detect_duplicate_rows(ws_original, header_row, col_map)
        duplicate_rows = {
            row: reason
            for row, reason in duplicate_rows.items()
            if row not in anulable_rows
        }

    pre_rows_coo = []
    pre_rows_other = []
    pedir_rows_coo = []
    pedir_rows_other = []

    total_rows = max(1, ws_original.max_row - header_row)
    emit("Procesando filas...", 0.25)

    for idx, r in enumerate(range(header_row + 1, ws_original.max_row + 1), start=1):
        if idx % 150 == 0:
            emit(f"Procesando filas... ({idx}/{total_rows})", 0.25 + 0.60 * (idx / total_rows))

        values = [ws_original.cell(r, c).value for c in range(1, ws_original.max_column + 1)]

        if r in anulable_rows:
            estado_val = ws_original.cell(r, col_estado).value if col_estado else ""
            estado_txt = norm(estado_val)
            reason = "Grupo anulable (neto saldo=0)"
            if estado_txt:
                reason += f" | Estado: {estado_txt}"
            ws_anul.append(values + [reason])
            continue

        if r in duplicate_rows:
            if "DUPLICADOS" in extra_sheets:
                extra_sheets["DUPLICADOS"].append(values + [duplicate_rows[r]])
            continue

        no_fac_val = ws_original.cell(r, col_no_fac).value
        contrato_val = ws_original.cell(r, col_contrato).value
        cod_proc_val = ws_original.cell(r, col_cod_proc).value
        proc_val = ws_original.cell(r, col_proc).value if col_proc else ""
        saldo_val = ws_original.cell(r, col_saldo).value if col_saldo else 0.0
        tipo_fac_val = ws_original.cell(r, col_tipo_fac).value if col_tipo_fac else ""
        id_paciente_val = ws_original.cell(r, col_id_paciente).value if col_id_paciente else ""
        nombre_paciente_val = ws_original.cell(r, col_nombre_paciente).value if col_nombre_paciente else ""

        # Saltar filas de totales o resumen del archivo original
        if (
            not norm(no_fac_val)
            and not norm(contrato_val)
            and not norm(cod_proc_val)
            and not norm(proc_val)
        ):
            continue

        ok, reason = specialty.apply_common_exclusion_rules(
            no_fac_val,
            contrato_val,
            cod_proc_val,
            tipo_fac_value=tipo_fac_val,
            id_paciente_value=id_paciente_val,
            nombre_paciente_value=nombre_paciente_val,
        )
        if not ok:
            ws_amar.append(values + [reason])
            continue

        row_kwargs = dict(
            row_values=values,
            contrato_value=contrato_val,
            procedimiento_value=proc_val,
            cod_proc_value=cod_proc_val,
            no_fac_value=no_fac_val,
            saldo_value=saldo_val,
            tarifario=tarifario,
            hosvi_by_fact=hosvi_by_fact,
            hosvi_by_alt=hosvi_by_alt,
            ws_original=ws_original,
            row_number=r,
            col_map=col_map,
            idx_vlr_aut=idx_vlr_aut,
            hosvi_cursor_fact=hosvi_cursor_fact,
            hosvi_cursor_alt=hosvi_cursor_alt,
        )

        if especialidad == "cardiologia":
            row_kwargs["payment_cfg"] = payment_cfg
        else:
            row_kwargs["payment_pct"] = payment_cfg.get("payment_pct")

        row_result = specialty.process_row(**row_kwargs)

        pre_values = list(values)
        pedir_values = list(values)

        if row_result.get("nuevo_vlr_aut") is not None and idx_vlr_aut is not None:
            pre_values[idx_vlr_aut - 1] = row_result["nuevo_vlr_aut"]
            pedir_values[idx_vlr_aut - 1] = row_result["nuevo_vlr_aut"]

        pre_row = pre_values + row_result["calc_values"]

        if row_result.get("is_coo_proced"):
            pre_rows_coo.append(pre_row)
            pedir_rows_coo.append(pedir_values)
        else:
            pre_rows_other.append(pre_row)
            pedir_rows_other.append(pedir_values)

    emit("Escribiendo hojas y totales...", 0.88)

    pre_rows_final = pre_rows_coo + pre_rows_other
    pedir_rows_final = pedir_rows_coo + pedir_rows_other
    duplicate_pre_rows = []
    if hasattr(specialty, "filter_duplicate_consult_rows"):
        (
            pre_rows_final,
            pedir_rows_final,
            duplicate_pre_rows,
        ) = specialty.filter_duplicate_consult_rows(
            pre_rows_coo=pre_rows_coo,
            pre_rows_other=pre_rows_other,
            pedir_rows_coo=pedir_rows_coo,
            pedir_rows_other=pedir_rows_other,
            base_headers=base_headers,
        )

    for row in pre_rows_final:
        ws_pre.append(row)

    ws_duplicates = extra_sheets.get("DUPLICADOS")
    if ws_duplicates is not None:
        for row in duplicate_pre_rows:
            ws_duplicates.append(row)

    for row in pedir_rows_final:
        ws_pedir.append(row)

    add_pre_summary_box(
        ws_pre,
        start_row=2,
        idx_saldo=idx_saldo,
        idx_vlr_aut=idx_vlr_aut,
        idx_total_coosalud=idx_total_coosalud,
    )
    add_pedir_fact_summary_box(
        ws_pedir,
        start_row=2,
        idx_vlr_aut=idx_vlr_aut,
        anchor_col=(idx_saldo - 1) if idx_saldo else None,
    )

    all_sheets = [ws_pre, ws_amar, ws_pedir, ws_anul] + list(extra_sheets.values())

    for ws in all_sheets:
        ws.freeze_panes = "A2"
        autofit_columns(ws)

    emit("Guardando archivo...", 0.96)
    out.save(output_path)
    emit("Listo.", 1.0)

    return ProcessResult(
        output_path=output_path,
        pre_rows=ws_pre.max_row - 1,
        pedir_rows=ws_pedir.max_row - 1,
        amarillo_rows=ws_amar.max_row - 1,
        anulados_rows=ws_anul.max_row - 1,
    )
