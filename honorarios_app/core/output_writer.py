from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
BOLD = Font(bold=True)
RED_BOLD = Font(bold=True, color="FF0000")
THIN = Side(style="thin", color="000000")
BOX_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws):
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(1, c)
        cell.fill = HEADER_FILL
        cell.font = BOLD


def autofit_columns(ws, max_width=60):
    for col in range(1, ws.max_column + 1):
        max_len = 0
        for row in range(1, min(ws.max_row, 2000) + 1):
            v = ws.cell(row, col).value
            if v is None:
                continue
            max_len = max(max_len, len(str(v)))
        ws.column_dimensions[get_column_letter(col)].width = min(max(10, max_len + 2), max_width)


def add_totals_row(ws, start_row, idx_saldo, idx_vlr_aut, label_col=1):
    last_data_row = ws.max_row
    totals_row = last_data_row + 2

    ws.cell(totals_row, label_col).value = "TOTALES"
    ws.cell(totals_row, label_col).font = BOLD

    if idx_saldo:
        c = get_column_letter(idx_saldo)
        ws.cell(totals_row, idx_saldo).value = f"=SUM({c}{start_row}:{c}{last_data_row})"
        ws.cell(totals_row, idx_saldo).font = BOLD

    if idx_vlr_aut:
        c = get_column_letter(idx_vlr_aut)
        ws.cell(totals_row, idx_vlr_aut).value = f"=SUM({c}{start_row}:{c}{last_data_row})"
        ws.cell(totals_row, idx_vlr_aut).font = BOLD


def _style_box_cell(cell, align="right"):
    cell.font = RED_BOLD
    cell.border = BOX_BORDER
    cell.alignment = Alignment(horizontal=align)


def add_pre_summary_box(ws, start_row, idx_saldo, idx_vlr_aut, idx_total_coosalud):
    if not idx_saldo or not idx_vlr_aut:
        return

    last_data_row = ws.max_row
    if last_data_row < start_row:
        return

    label_col = max(1, idx_saldo - 1)
    row_1 = last_data_row + 2
    row_2 = row_1 + 2
    row_3 = row_2 + 1

    saldo_col_letter = get_column_letter(idx_saldo)
    vlr_col_letter = get_column_letter(idx_vlr_aut)

    # TOTAL BAJAR HOSV
    c = ws.cell(row_1, label_col)
    c.value = "TOTAL BAJAR HOSV"
    _style_box_cell(c)

    c = ws.cell(row_1, idx_saldo)
    c.value = f"=SUM({saldo_col_letter}{start_row}:{saldo_col_letter}{last_data_row})"
    _style_box_cell(c)

    c = ws.cell(row_1, idx_vlr_aut)
    c.value = f"=SUM({vlr_col_letter}{start_row}:{vlr_col_letter}{last_data_row})"
    _style_box_cell(c)

    # TOTAL COOSALUD
    # Debe ser el mismo valor de la suma de VLR. A AUTORIZAR
    c = ws.cell(row_2, label_col)
    c.value = "TOTAL COOSALUD"
    _style_box_cell(c)

    c = ws.cell(row_2, idx_saldo)
    c.value = f"=SUM({vlr_col_letter}{start_row}:{vlr_col_letter}{last_data_row})"
    _style_box_cell(c)

    # DIF = TOTAL BAJAR HOSV - TOTAL COOSALUD
    c = ws.cell(row_3, label_col)
    c.value = "DIF"
    _style_box_cell(c)

    c = ws.cell(row_3, idx_saldo)
    c.value = f"={saldo_col_letter}{row_1}-{saldo_col_letter}{row_2}"
    _style_box_cell(c)


def add_pedir_fact_summary_box(ws, start_row, idx_vlr_aut, anchor_col=None, label_text="TOTAL PEDIR FACT"):
    if not idx_vlr_aut:
        return

    last_data_row = ws.max_row
    if last_data_row < start_row:
        return

    label_col = anchor_col if anchor_col is not None else max(1, idx_vlr_aut - 2)
    summary_row = last_data_row + 2
    vlr_col_letter = get_column_letter(idx_vlr_aut)

    c = ws.cell(summary_row, label_col)
    c.value = label_text
    c.font = RED_BOLD
    c.alignment = Alignment(horizontal="right")

    c = ws.cell(summary_row, idx_vlr_aut)
    c.value = f"=SUM({vlr_col_letter}{start_row}:{vlr_col_letter}{last_data_row})"
    _style_box_cell(c)