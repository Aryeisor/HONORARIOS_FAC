from openpyxl import load_workbook

from honorarios_app.core.common import (
    find_header_row,
    get_col,
    norm,
    normalize_datetime_key,
    normalize_fac,
    parse_percent,
    to_number,
)


def validate_required_sheets(
    input_path: str,
    original_sheet=None,
    hosvi_sheet: str = "HOSVIREPORT",
    coosalud_sheet: str = "COOSALUD",
) -> None:
    wb = load_workbook(input_path, read_only=True, data_only=True)

    missing = []
    if hosvi_sheet not in wb.sheetnames:
        missing.append(hosvi_sheet)
    if coosalud_sheet not in wb.sheetnames:
        missing.append(coosalud_sheet)

    if missing:
        raise RuntimeError("Falta(n) hoja(s): " + ", ".join(missing))


def load_excel(input_path):
    return load_workbook(input_path, data_only=True)


def build_tarifario_coosalud(ws):
    header_row, col_map = find_header_row(ws, ["CIRUJ", "AYUD"])
    if header_row is None:
        raise RuntimeError("No pude detectar encabezados en la hoja COOSALUD (busqué CIRUJ / AYUD).")

    col_cod = get_col(col_map, "CUPS", "CODIGO", "CODIGO PROCEDIMIENTO", "CÓDIGO", "COD", "UNNAMED: 0")
    if not col_cod:
        col_cod = 1

    col_cir = (
        get_col(col_map, "CIRUJANO")
        or get_col(col_map, "CIRUJANO CONTRATADO CON COOSALUD")
        or get_col(col_map, "100% CIR", "100% CIRUJANO")
    )

    col_ayd2 = (
        get_col(col_map, "AYUDANTE 2")
        or get_col(col_map, "AYUDANTE 2 CONTRATADO CON COOSALUD")
        or get_col(col_map, "100% AYUD", "AYUD.2", "AYUDANTE2")
    )

    if not col_cir or not col_ayd2:
        raise RuntimeError("Faltan columnas de tarifa en COOSALUD (CIRUJANO y AYUDANTE 2).")

    tarifario = {}
    for r in range(header_row + 1, ws.max_row + 1):
        cod = normalize_fac(ws.cell(r, col_cod).value)
        if not cod:
            continue

        tarifario[cod] = {
            "cir": to_number(ws.cell(r, col_cir).value),
            "ayd2": to_number(ws.cell(r, col_ayd2).value),
        }

    return tarifario


def build_pct_override_if_present(wb, sheet_name="PORCENTAJE_PAGO"):
    if sheet_name not in wb.sheetnames:
        return {}

    ws = wb[sheet_name]
    header_row, col_map = find_header_row(ws, ["FACT", "COD", "PORC"])
    if header_row is None:
        return {}

    col_fact = get_col(col_map, "FACTURA", "NO FAC", "N° FAC")
    col_cod = get_col(col_map, "COD_PROCED", "CODIGO PROCEDIMIENTO", "CODIGO", "COD")
    col_pct = get_col(col_map, "PORC_PAGO", "PORC PAGO", "PORCENTAJE", "%", "PORC")

    if not col_fact or not col_cod or not col_pct:
        return {}

    out = {}
    for r in range(header_row + 1, ws.max_row + 1):
        f = normalize_fac(ws.cell(r, col_fact).value)
        c = normalize_fac(ws.cell(r, col_cod).value)
        p = parse_percent(ws.cell(r, col_pct).value)

        if f and c and p is not None:
            out[(f, c)] = p

    return out


def build_hosvireport(ws, pct_override=None):
    """
    Devuelve dos índices:
    1) by_fact: (factura, cod_proced) -> lista de registros
    2) by_alt:  (documento, cod_proced, fecha_proced) -> lista de registros
    """
    pct_override = pct_override or {}

    header_row, col_map = find_header_row(ws, ["FACTURA", "COD_PROCED"])
    if header_row is None:
        raise RuntimeError("No pude detectar encabezados en HOSVIREPORT (busqué FACTURA y COD_PROCED).")

    col_fact = get_col(col_map, "FACTURA")
    col_codp = get_col(col_map, "COD_PROCED", "COD PROCED", "CODIGO PROCED", "CODIGO PROCEDIMIENTO")
    col_rol = get_col(col_map, "DES_HONORARIO", "HONORARIO", "DES HONORARIO")
    col_pct = get_col(col_map, "%_FACTURADO", "% FACTURADO", "FACTURADO", "PORC")
    col_doc = get_col(col_map, "DOCUMENTO", "IDENTIFICACION", "DOC")
    col_fecha = get_col(col_map, "FECHA_PROCED", "FECHA PROCED", "FECHA DEL SERVICIO", "FECHA SERVICIO", "FECHA")
    col_est = get_col(col_map, "ESTADO")

    if not col_codp:
        raise RuntimeError("No encontré COD_PROCED en HOSVIREPORT.")

    by_fact = {}
    by_alt = {}

    for r in range(header_row + 1, ws.max_row + 1):
        factura = normalize_fac(ws.cell(r, col_fact).value) if col_fact else ""
        codp = normalize_fac(ws.cell(r, col_codp).value)
        doc = normalize_fac(ws.cell(r, col_doc).value) if col_doc else ""
        fecha = normalize_datetime_key(ws.cell(r, col_fecha).value) if col_fecha else ""

        if not codp:
            continue

        if col_est:
            est = norm(ws.cell(r, col_est).value)
            if est == "ANULADO":
                continue

        rol = norm(ws.cell(r, col_rol).value) if col_rol else ""
        pct = parse_percent(ws.cell(r, col_pct).value) if col_pct else None

        row_data = {"rol": rol, "pct": pct}

        if factura:
            key_fact = (factura, codp)
            if key_fact in pct_override:
                row_data["pct"] = pct_override[key_fact]
            by_fact.setdefault(key_fact, []).append(row_data.copy())

        if doc and fecha:
            key_alt = (doc, codp, fecha)
            by_alt.setdefault(key_alt, []).append(row_data.copy())

    return by_fact, by_alt