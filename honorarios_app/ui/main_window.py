import os
import sys
import subprocess
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from datetime import datetime

from honorarios_app.config.settings import get_specialty_settings, set_specialty_settings
from honorarios_app.core.processor import process_excel


APP_TITLE = "Honorarios Médicos Por Facturado"
APP_SUBTITLE = "Automatización de liquidación de honorarios por especialidad"


class MainWindow:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry("1360x810")
        self.root.minsize(1000, 620)
        self.root.configure(bg="#EEF7F5")

        try:
            self.root.state("zoomed")
        except Exception:
            pass

        self.input_var = tk.StringVar()
        self.output_var = tk.StringVar()
        self.specialty_var = tk.StringVar(value="ortopedia")
        self.status_var = tk.StringVar(value="Listo para procesar.")
        self.progress_var = tk.DoubleVar(value=0)

        # Ortopedia / Urología / Fonoaudiología
        self.base_pct_var = tk.StringVar(value="70")
        self.update_pct_var = tk.StringVar(value="70")

        # Cardiología
        self.base_consulta_pct_var = tk.StringVar(value="90")
        self.update_consulta_pct_var = tk.StringVar(value="90")
        self.base_proced_pct_var = tk.StringVar(value="36")
        self.update_proced_pct_var = tk.StringVar(value="36")

        self.output_auto = True
        self.is_processing = False
        self.last_result = None

        self.main_canvas = None
        self.scrollable_frame = None
        self.canvas_window = None
        self.content = None
        self.form_card = None
        self.actions_card = None
        self.log_card = None

        self._configure_styles()
        self._build_ui()
        self._bind_events()
        self._load_specialty_percentage()
        self._toggle_percentage_mode()
        self._on_root_resize()

    def _configure_styles(self):
        style = ttk.Style()

        try:
            style.theme_use("clam")
        except Exception:
            pass

        style.configure("App.TFrame", background="#EEF7F5")
        style.configure("White.TFrame", background="#FFFFFF")
        style.configure("Panel.TFrame", background="#F4FBF9")

        style.configure(
            "Title.TLabel",
            background="#EEF7F5",
            foreground="#0F172A",
            font=("Segoe UI", 22, "bold"),
        )
        style.configure(
            "Subtitle.TLabel",
            background="#EEF7F5",
            foreground="#4B5563",
            font=("Segoe UI", 10),
        )
        style.configure(
            "CardTitle.TLabel",
            background="#FFFFFF",
            foreground="#0F172A",
            font=("Segoe UI", 11, "bold"),
        )
        style.configure(
            "CardText.TLabel",
            background="#FFFFFF",
            foreground="#475569",
            font=("Segoe UI", 10),
        )
        style.configure(
            "Field.TLabel",
            background="#FFFFFF",
            foreground="#334155",
            font=("Segoe UI", 10),
        )
        style.configure(
            "MetricValue.TLabel",
            background="#FFFFFF",
            foreground="#111827",
            font=("Segoe UI", 18, "bold"),
        )
        style.configure(
            "MetricLabel.TLabel",
            background="#FFFFFF",
            foreground="#64748B",
            font=("Segoe UI", 9),
        )
        style.configure(
            "Primary.TButton",
            font=("Segoe UI", 10, "bold"),
            padding=10,
            background="#7BC4B8",
            foreground="#083344",
        )
        style.map(
            "Primary.TButton",
            background=[("active", "#6BB7AA"), ("pressed", "#5FA99D")],
            foreground=[("active", "#083344")],
        )

        style.configure(
            "Secondary.TButton",
            font=("Segoe UI", 10),
            padding=8,
            background="#DDF3EF",
            foreground="#0F172A",
        )
        style.map(
            "Secondary.TButton",
            background=[("active", "#C9ECE6"), ("pressed", "#B8E5DE")],
            foreground=[("active", "#0F172A")],
        )

        style.configure(
            "Status.TLabel",
            background="#FFFFFF",
            foreground="#334155",
            font=("Segoe UI", 10),
        )
        style.configure(
            "Hint.TLabel",
            background="#F4FBF9",
            foreground="#0F766E",
            font=("Segoe UI", 9),
            padding=10,
        )
        style.configure("TEntry", padding=7)
        style.configure("TCombobox", padding=5)
        style.configure("Horizontal.TProgressbar", thickness=16)

    def _build_ui(self):
        outer = ttk.Frame(self.root, style="App.TFrame")
        outer.pack(fill="both", expand=True)

        self.main_canvas = tk.Canvas(
            outer,
            bg="#EEF7F5",
            highlightthickness=0,
            bd=0
        )
        self.main_canvas.pack(side="left", fill="both", expand=True)

        self.v_scroll = ttk.Scrollbar(outer, orient="vertical", command=self.main_canvas.yview)
        self.v_scroll.pack(side="right", fill="y")

        self.main_canvas.configure(yscrollcommand=self.v_scroll.set)

        self.scrollable_frame = ttk.Frame(self.main_canvas, style="App.TFrame", padding=18)
        self.canvas_window = self.main_canvas.create_window(
            (0, 0),
            window=self.scrollable_frame,
            anchor="nw"
        )

        self.scrollable_frame.bind("<Configure>", self._on_frame_configure)
        self.main_canvas.bind("<Configure>", self._on_canvas_configure)

        self._build_header(self.scrollable_frame)
        self._build_metrics(self.scrollable_frame)

        self.content = ttk.Frame(self.scrollable_frame, style="App.TFrame")
        self.content.pack(fill="both", expand=True, pady=(14, 0))

        self.content.columnconfigure(0, weight=3)
        self.content.columnconfigure(1, weight=1)
        self.content.rowconfigure(2, weight=1)

        self._build_form_card(self.content)
        self._build_actions_card(self.content)
        self._build_log_card(self.content)

        self.root.bind("<Configure>", self._on_root_resize)
        self._bind_mousewheel()

    def _on_frame_configure(self, event=None):
        if self.main_canvas is not None:
            self.main_canvas.configure(scrollregion=self.main_canvas.bbox("all"))

    def _on_canvas_configure(self, event):
        if self.main_canvas is not None and self.canvas_window is not None:
            self.main_canvas.itemconfig(self.canvas_window, width=event.width)

    def _bind_mousewheel(self):
        def _on_mousewheel(event):
            if self.main_canvas is not None:
                self.main_canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

        def _on_mousewheel_linux_up(event):
            if self.main_canvas is not None:
                self.main_canvas.yview_scroll(-1, "units")

        def _on_mousewheel_linux_down(event):
            if self.main_canvas is not None:
                self.main_canvas.yview_scroll(1, "units")

        self.main_canvas.bind_all("<MouseWheel>", _on_mousewheel)
        self.main_canvas.bind_all("<Button-4>", _on_mousewheel_linux_up)
        self.main_canvas.bind_all("<Button-5>", _on_mousewheel_linux_down)

    def _on_root_resize(self, event=None):
        if self.content is None or self.form_card is None or self.actions_card is None or self.log_card is None:
            return

        width = self.root.winfo_width()

        if width < 1280:
            self.content.columnconfigure(0, weight=1)
            self.content.columnconfigure(1, weight=0)

            self.form_card.grid_configure(row=0, column=0, padx=(0, 0), pady=(0, 12), sticky="nsew")
            self.actions_card.grid_configure(row=1, column=0, padx=(0, 0), pady=(0, 12), sticky="ew")
            self.log_card.grid_configure(row=2, column=0, columnspan=1, pady=(0, 0), sticky="nsew")
        else:
            self.content.columnconfigure(0, weight=3)
            self.content.columnconfigure(1, weight=1)

            self.form_card.grid_configure(row=0, column=0, padx=(0, 12), pady=(0, 0), sticky="nsew")
            self.actions_card.grid_configure(row=0, column=1, padx=(0, 0), pady=(0, 0), sticky="nsew")
            self.log_card.grid_configure(row=1, column=0, columnspan=2, pady=(12, 0), sticky="nsew")

        self.root.after_idle(self._on_frame_configure)

    def _build_header(self, parent):
        header = ttk.Frame(parent, style="App.TFrame")
        header.pack(fill="x")

        ttk.Label(header, text=APP_TITLE, style="Title.TLabel").pack(anchor="w")
        ttk.Label(header, text=APP_SUBTITLE, style="Subtitle.TLabel").pack(anchor="w", pady=(2, 0))

    def _build_metrics(self, parent):
        metrics = ttk.Frame(parent, style="App.TFrame")
        metrics.pack(fill="x", pady=(16, 0))

        metrics.columnconfigure((0, 1, 2, 3), weight=1)

        self.metric_specialty = self._create_metric_card(metrics, 0, "Especialidad", "ORTOPEDIA")
        self.metric_pre = self._create_metric_card(metrics, 1, "Registros PRE", "0")
        self.metric_pedir = self._create_metric_card(metrics, 2, "Registros PEDIR FAC", "0")
        self.metric_estado = self._create_metric_card(metrics, 3, "Estado del módulo", "LISTO")

    def _create_metric_card(self, parent, column, label, value):
        card = ttk.Frame(parent, style="White.TFrame", padding=14)
        card.grid(row=0, column=column, sticky="nsew", padx=(0 if column == 0 else 8, 0))

        value_lbl = ttk.Label(card, text=value, style="MetricValue.TLabel")
        value_lbl.pack(anchor="w")

        ttk.Label(card, text=label, style="MetricLabel.TLabel").pack(anchor="w", pady=(4, 0))
        return value_lbl

    def _build_form_card(self, parent):
        self.form_card = ttk.Frame(parent, style="White.TFrame", padding=18)
        self.form_card.grid(row=0, column=0, sticky="nsew", padx=(0, 12))
        card = self.form_card

        card.columnconfigure(1, weight=1)
        card.columnconfigure(2, weight=1)

        ttk.Label(card, text="Parámetros de procesamiento", style="CardTitle.TLabel").grid(
            row=0, column=0, columnspan=4, sticky="w", pady=(0, 14)
        )

        ttk.Label(card, text="Especialidad", style="Field.TLabel").grid(row=1, column=0, sticky="w", pady=6)

        top_row_frame = ttk.Frame(card, style="White.TFrame")
        top_row_frame.grid(row=1, column=1, columnspan=3, sticky="ew", pady=6)

        self.combo_specialty = ttk.Combobox(
            top_row_frame,
            textvariable=self.specialty_var,
            values=["ortopedia", "urologia", "cardiologia", "fonoaudiologia"],
            state="readonly",
            width=22,
        )
        self.combo_specialty.grid(row=0, column=0, sticky="w", padx=(0, 14))

        self.simple_pct_frame = ttk.Frame(top_row_frame, style="White.TFrame")
        self.simple_pct_frame.grid(row=0, column=1, sticky="w")

        ttk.Label(self.simple_pct_frame, text="Valor % base", style="Field.TLabel").grid(row=0, column=0, sticky="w", padx=(0, 6))
        self.entry_base_pct = ttk.Entry(self.simple_pct_frame, textvariable=self.base_pct_var, width=8, state="readonly")
        self.entry_base_pct.grid(row=0, column=1, sticky="w", padx=(0, 14))

        ttk.Label(self.simple_pct_frame, text="Valor % a actualizar", style="Field.TLabel").grid(row=0, column=2, sticky="w", padx=(0, 6))
        self.entry_update_pct = ttk.Entry(self.simple_pct_frame, textvariable=self.update_pct_var, width=8)
        self.entry_update_pct.grid(row=0, column=3, sticky="w")

        self.cardio_pct_frame = ttk.Frame(top_row_frame, style="White.TFrame")
        self.cardio_pct_frame.grid(row=0, column=1, sticky="w")

        ttk.Label(self.cardio_pct_frame, text="% base consultas/cuidados", style="Field.TLabel").grid(row=0, column=0, sticky="w", padx=(0, 6))
        self.entry_base_consulta_pct = ttk.Entry(
            self.cardio_pct_frame, textvariable=self.base_consulta_pct_var, width=8, state="readonly"
        )
        self.entry_base_consulta_pct.grid(row=0, column=1, sticky="w", padx=(0, 12))

        ttk.Label(self.cardio_pct_frame, text="% actualizar consultas/cuidados", style="Field.TLabel").grid(row=0, column=2, sticky="w", padx=(0, 6))
        self.entry_update_consulta_pct = ttk.Entry(
            self.cardio_pct_frame, textvariable=self.update_consulta_pct_var, width=8
        )
        self.entry_update_consulta_pct.grid(row=0, column=3, sticky="w", padx=(0, 18))

        ttk.Label(self.cardio_pct_frame, text="% base procedimientos", style="Field.TLabel").grid(row=1, column=0, sticky="w", padx=(0, 6), pady=(8, 0))
        self.entry_base_proced_pct = ttk.Entry(
            self.cardio_pct_frame, textvariable=self.base_proced_pct_var, width=8, state="readonly"
        )
        self.entry_base_proced_pct.grid(row=1, column=1, sticky="w", padx=(0, 12), pady=(8, 0))

        ttk.Label(self.cardio_pct_frame, text="% actualizar procedimientos", style="Field.TLabel").grid(row=1, column=2, sticky="w", padx=(0, 6), pady=(8, 0))
        self.entry_update_proced_pct = ttk.Entry(
            self.cardio_pct_frame, textvariable=self.update_proced_pct_var, width=8
        )
        self.entry_update_proced_pct.grid(row=1, column=3, sticky="w", pady=(8, 0))

        ttk.Label(card, text="Archivo de entrada", style="Field.TLabel").grid(row=2, column=0, sticky="w", pady=6)
        self.entry_input = ttk.Entry(card, textvariable=self.input_var)
        self.entry_input.grid(row=2, column=1, columnspan=2, sticky="ew", pady=6)
        self.btn_input = ttk.Button(card, text="Buscar archivo", style="Secondary.TButton", command=self.select_input)
        self.btn_input.grid(row=2, column=3, padx=(8, 0), pady=6)

        ttk.Label(card, text="Archivo de salida", style="Field.TLabel").grid(row=3, column=0, sticky="w", pady=6)
        self.entry_output = ttk.Entry(card, textvariable=self.output_var)
        self.entry_output.grid(row=3, column=1, columnspan=2, sticky="ew", pady=6)
        self.btn_output = ttk.Button(card, text="Guardar como", style="Secondary.TButton", command=self.select_output)
        self.btn_output.grid(row=3, column=3, padx=(8, 0), pady=6)

        hint = (
            "El sistema sugiere automáticamente el nombre del archivo de salida con el formato "
            "'PEDIR FAC + especialidad + nombre del archivo de entrada'. "
            "Los porcentajes base pueden modificarse temporalmente o guardarse como nuevos valores base."
        )
        ttk.Label(card, text=hint, style="Hint.TLabel", wraplength=980, justify="left").grid(
            row=4, column=0, columnspan=4, sticky="ew", pady=(14, 14)
        )

        buttons = ttk.Frame(card, style="White.TFrame")
        buttons.grid(row=5, column=0, columnspan=4, sticky="ew")
        buttons.columnconfigure((0, 1, 2), weight=1)

        self.btn_process = ttk.Button(
            buttons,
            text="Procesar archivo",
            style="Primary.TButton",
            command=self.process,
        )
        self.btn_process.grid(row=0, column=0, sticky="ew", padx=(0, 8))

        self.btn_clear = ttk.Button(
            buttons,
            text="Limpiar formulario",
            style="Secondary.TButton",
            command=self.clear_fields,
        )
        self.btn_clear.grid(row=0, column=1, sticky="ew", padx=4)

        self.btn_open_folder = ttk.Button(
            buttons,
            text="Abrir ubicación del archivo",
            style="Secondary.TButton",
            command=self.open_output_folder,
        )
        self.btn_open_folder.grid(row=0, column=2, sticky="ew", padx=(8, 0))

        progress_frame = ttk.Frame(card, style="White.TFrame")
        progress_frame.grid(row=6, column=0, columnspan=4, sticky="ew", pady=(18, 0))
        progress_frame.columnconfigure(0, weight=1)

        self.progress = ttk.Progressbar(
            progress_frame,
            style="Horizontal.TProgressbar",
            variable=self.progress_var,
            maximum=100,
        )
        self.progress.grid(row=0, column=0, sticky="ew")

        ttk.Label(progress_frame, textvariable=self.status_var, style="Status.TLabel").grid(
            row=1, column=0, sticky="w", pady=(8, 0)
        )

    def _build_actions_card(self, parent):
        self.actions_card = ttk.Frame(parent, style="White.TFrame", padding=18)
        self.actions_card.grid(row=0, column=1, sticky="nsew")
        card = self.actions_card

        ttk.Label(card, text="Soporte y consulta", style="CardTitle.TLabel").pack(anchor="w", pady=(0, 12))

        self.btn_help = ttk.Button(card, text="Ayuda de uso", style="Secondary.TButton", command=self.show_help)
        self.btn_help.pack(fill="x", pady=4)

        self.btn_info = ttk.Button(card, text="Guía del módulo", style="Secondary.TButton", command=self.show_info)
        self.btn_info.pack(fill="x", pady=4)

        self.btn_manual = ttk.Button(card, text="Manual del módulo", style="Secondary.TButton", command=self.open_manual_pdf)
        self.btn_manual.pack(fill="x", pady=4)

        self.btn_about = ttk.Button(card, text="Acerca del sistema", style="Secondary.TButton", command=self.show_about)
        self.btn_about.pack(fill="x", pady=4)

        self.btn_log = ttk.Button(card, text="Ver bitácora", style="Secondary.TButton", command=self.focus_log)
        self.btn_log.pack(fill="x", pady=4)

        self.btn_report = ttk.Button(card, text="Resumen del proceso", style="Secondary.TButton", command=self.quick_report)
        self.btn_report.pack(fill="x", pady=4)

        ttk.Separator(card, orient="horizontal").pack(fill="x", pady=14)

        ttk.Label(card, text="Herramientas", style="CardTitle.TLabel").pack(anchor="w", pady=(0, 8))

        self.btn_validate = ttk.Button(card, text="Validar archivo", style="Secondary.TButton", command=self.validate_file)
        self.btn_validate.pack(fill="x", pady=4)

        self.btn_open_last = ttk.Button(card, text="Abrir último generado", style="Secondary.TButton", command=self.open_last_generated)
        self.btn_open_last.pack(fill="x", pady=4)

        self.btn_reset_pct = ttk.Button(card, text="Restablecer porcentajes", style="Secondary.TButton", command=self.reset_percentages)
        self.btn_reset_pct.pack(fill="x", pady=4)

        self.btn_copy_summary = ttk.Button(card, text="Copiar resumen", style="Secondary.TButton", command=self.copy_summary)
        self.btn_copy_summary.pack(fill="x", pady=4)

    def _build_log_card(self, parent):
        self.log_card = ttk.Frame(parent, style="White.TFrame", padding=18)
        self.log_card.grid(row=1, column=0, columnspan=2, sticky="nsew", pady=(12, 0))
        card = self.log_card

        card.rowconfigure(1, weight=1)
        card.columnconfigure(0, weight=1)

        ttk.Label(card, text="Bitácora del sistema", style="CardTitle.TLabel").grid(
            row=0, column=0, sticky="w", pady=(0, 10)
        )

        text_frame = ttk.Frame(card, style="White.TFrame")
        text_frame.grid(row=1, column=0, sticky="nsew")

        text_frame.rowconfigure(0, weight=1)
        text_frame.columnconfigure(0, weight=1)

        self.log_text = tk.Text(
            text_frame,
            wrap="word",
            font=("Consolas", 10),
            bg="#F4FBF9",
            fg="#0F172A",
            relief="solid",
            borderwidth=1,
            height=12,
        )
        self.log_text.grid(row=0, column=0, sticky="nsew")

        scroll = ttk.Scrollbar(text_frame, orient="vertical", command=self.log_text.yview)
        scroll.grid(row=0, column=1, sticky="ns")
        self.log_text.configure(yscrollcommand=scroll.set)

        self.log("Sistema iniciado correctamente.")
        self.log("Módulo de facturados disponible para procesamiento.")
        self._refresh_metrics()

    def _bind_events(self):
        self.combo_specialty.bind("<<ComboboxSelected>>", self.on_specialty_change)
        self.output_var.trace_add("write", self.on_output_change)

    def _toggle_percentage_mode(self):
        specialty = self.specialty_var.get().strip().lower()

        if specialty == "cardiologia":
            self.simple_pct_frame.grid_remove()
            self.cardio_pct_frame.grid()
        else:
            self.cardio_pct_frame.grid_remove()
            self.simple_pct_frame.grid()

    def _load_specialty_percentage(self):
        specialty = self.specialty_var.get().strip().lower()
        settings = get_specialty_settings(specialty)

        if specialty == "cardiologia":
            consulta_pct = settings.get("consulta_pct", 0.90)
            proced_pct = settings.get("procedimiento_pct", 0.36)

            self.base_consulta_pct_var.set(str(int(round(consulta_pct * 100))))
            self.update_consulta_pct_var.set(str(int(round(consulta_pct * 100))))
            self.base_proced_pct_var.set(str(int(round(proced_pct * 100))))
            self.update_proced_pct_var.set(str(int(round(proced_pct * 100))))
        else:
            payment_pct = settings.get("payment_pct", 0.70)
            pct_display = str(int(round(payment_pct * 100)))
            self.base_pct_var.set(pct_display)
            self.update_pct_var.set(pct_display)

    def _get_required_sheets_for_specialty(self, specialty: str):
        specialty = (specialty or "").strip().lower()

        if specialty in ("ortopedia", "urologia"):
            return ["COOSALUD", "HOSVIREPORT"]

        if specialty in ("cardiologia", "fonoaudiologia"):
            return ["COOSALUD"]

        return []

    def _get_resource_path(self, relative_path: str):
        if getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS"):
            base_path = sys._MEIPASS
        else:
            current_dir = os.path.dirname(os.path.abspath(__file__))
            base_path = os.path.dirname(current_dir)

        return os.path.join(base_path, relative_path)

    def on_specialty_change(self, event=None):
        self._refresh_metrics()
        self._load_specialty_percentage()
        self._toggle_percentage_mode()

        if self.input_var.get().strip() and self.output_auto:
            self.output_var.set(self._suggest_output_path(self.input_var.get().strip()))
            self.log("Nombre de salida actualizado automáticamente por cambio de especialidad.")

    def on_output_change(self, *args):
        current = self.output_var.get().strip()
        if not current:
            self.output_auto = True

    def _refresh_metrics(self):
        self.metric_specialty.config(text=self.specialty_var.get().strip().upper())
        if self.is_processing:
            self.metric_estado.config(text="PROCESANDO")
        else:
            self.metric_estado.config(text="LISTO")

    def _get_percent_value(self, raw_value, field_name):
        raw = raw_value.strip().replace("%", "").replace(",", ".")
        if not raw:
            raise ValueError(f"Debes ingresar {field_name}.")

        value = float(raw)
        if value <= 0 or value > 100:
            raise ValueError(f"{field_name} debe estar entre 1 y 100.")

        return value / 100.0

    def _build_payment_payload(self):
        specialty = self.specialty_var.get().strip().lower()

        if specialty == "cardiologia":
            consulta_pct = self._get_percent_value(
                self.update_consulta_pct_var.get(),
                "el porcentaje de consultas/cuidados",
            )
            procedimiento_pct = self._get_percent_value(
                self.update_proced_pct_var.get(),
                "el porcentaje de procedimientos",
            )

            return {
                "mode": "cardio",
                "payment_cfg": {
                    "consulta_pct": consulta_pct,
                    "procedimiento_pct": procedimiento_pct,
                }
            }

        payment_pct = self._get_percent_value(
            self.update_pct_var.get(),
            "el porcentaje a actualizar",
        )
        return {
            "mode": "simple",
            "payment_pct": payment_pct,
        }

    def log(self, message):
        timestamp = datetime.now().strftime("%H:%M:%S")
        self.log_text.insert("end", f"[{timestamp}] {message}\n")
        self.log_text.see("end")
        self.root.update_idletasks()
        self._on_frame_configure()

    def set_status(self, message, percent=None):
        self.status_var.set(message)
        if percent is not None:
            self.progress_var.set(percent)
        self.root.update_idletasks()

    def select_input(self):
        path = filedialog.askopenfilename(
            title="Seleccionar archivo Excel de entrada",
            filetypes=[("Archivos Excel", "*.xlsx *.xlsm *.xltx *.xltm")],
        )
        if path:
            self.input_var.set(path)
            self.log(f"Archivo de entrada seleccionado: {path}")

            if self.output_auto or not self.output_var.get().strip():
                self.output_var.set(self._suggest_output_path(path))
                self.output_auto = True
                self.log("Se generó automáticamente la ruta de salida sugerida.")

    def select_output(self):
        path = filedialog.asksaveasfilename(
            title="Definir archivo de salida",
            defaultextension=".xlsx",
            filetypes=[("Archivos Excel", "*.xlsx")],
        )
        if path:
            self.output_var.set(path)
            self.output_auto = False
            self.log(f"Archivo de salida definido manualmente: {path}")

    def _suggest_output_path(self, input_path):
        folder = os.path.dirname(input_path)
        specialty = self.specialty_var.get().strip().upper()
        base_name = os.path.basename(input_path)
        name_without_ext, _ = os.path.splitext(base_name)
        filename = f"PEDIR FAC_{specialty}_{name_without_ext}.xlsx"
        return os.path.join(folder, filename)

    def clear_fields(self):
        self.input_var.set("")
        self.output_var.set("")
        self.specialty_var.set("ortopedia")
        self.progress_var.set(0)
        self.status_var.set("Listo para procesar.")
        self.output_auto = True
        self.last_result = None
        self.metric_pre.config(text="0")
        self.metric_pedir.config(text="0")
        self.metric_estado.config(text="LISTO")
        self.metric_specialty.config(text="ORTOPEDIA")
        self._load_specialty_percentage()
        self._toggle_percentage_mode()
        self.log("El formulario fue restablecido correctamente.")

    def open_output_folder(self):
        output_path = self.last_result.output_path if self.last_result and getattr(self.last_result, "output_path", None) else self.output_var.get().strip()

        if not output_path:
            messagebox.showwarning(
                "Ubicación no disponible",
                "Todavía no se ha definido una ruta de salida válida.\n\nSeleccione un archivo de salida o procese un archivo primero."
            )
            return

        output_path = os.path.normpath(output_path)
        folder = os.path.dirname(output_path)

        if not folder or not os.path.exists(folder):
            messagebox.showwarning(
                "Ubicación no encontrada",
                "La carpeta asociada al archivo de salida no existe o no está disponible en este momento."
            )
            return

        try:
            if os.path.exists(output_path):
                self.log(f"Abriendo explorador y seleccionando archivo: {output_path}")
                subprocess.Popen(["explorer", "/select,", output_path])
            else:
                os.startfile(folder)
                messagebox.showinfo(
                    "Archivo aún no generado",
                    "Se abrió la carpeta de salida, pero el archivo final todavía no existe.\n\nProcese el archivo para generar el resultado."
                )
        except Exception as e:
            messagebox.showerror(
                "No fue posible abrir la ubicación",
                f"Ocurrió un inconveniente al intentar abrir la carpeta o seleccionar el archivo.\n\nDetalle técnico:\n{e}"
            )

    def open_last_generated(self):
        if not self.last_result or not getattr(self.last_result, "output_path", None):
            messagebox.showinfo(
                "Archivo no disponible",
                "Todavía no existe un archivo generado en esta sesión.\n\nProcese un archivo y luego utilice esta opción para abrir el resultado directamente."
            )
            return

        output_path = os.path.normpath(self.last_result.output_path)

        if not os.path.exists(output_path):
            messagebox.showwarning(
                "Archivo no encontrado",
                "El último archivo registrado no se encuentra disponible en la ubicación esperada."
            )
            return

        try:
            os.startfile(output_path)
            self.log(f"Se abrió el último archivo generado: {output_path}")
        except Exception as e:
            messagebox.showerror(
                "No fue posible abrir el archivo",
                f"Ocurrió un inconveniente al intentar abrir el último archivo generado.\n\nDetalle técnico:\n{e}"
            )

    def open_manual_pdf(self):
        try:
            manual_path = self._get_resource_path(
                os.path.join("resources", "manual_usuario_facturados.pdf")
            )

            if not os.path.exists(manual_path):
                messagebox.showerror(
                    "Manual no encontrado",
                    f"No se encontró el manual en la ruta esperada:\n\n{manual_path}"
                )
                return

            os.startfile(manual_path)
            self.log("Manual de usuario abierto correctamente.")

        except Exception as e:
            messagebox.showerror(
                "No fue posible abrir el manual",
                f"Ocurrió un inconveniente al intentar abrir el manual de usuario.\n\nDetalle técnico:\n{e}"
            )

    def update_progress(self, msg, frac):
        percent = int(max(0.0, min(1.0, frac)) * 100)
        self.set_status(f"{msg} ({percent}%)", percent)
        self.log(msg)
        self.metric_estado.config(text="PROCESANDO")

    def set_processing_state(self, processing: bool):
        self.is_processing = processing
        state = "disabled" if processing else "normal"

        self.combo_specialty.configure(state="disabled" if processing else "readonly")
        self.entry_input.configure(state=state)
        self.entry_output.configure(state=state)

        self.entry_update_pct.configure(state=state)
        self.entry_base_pct.configure(state="readonly")

        self.entry_update_consulta_pct.configure(state=state)
        self.entry_base_consulta_pct.configure(state="readonly")
        self.entry_update_proced_pct.configure(state=state)
        self.entry_base_proced_pct.configure(state="readonly")

        for btn in (
            self.btn_input,
            self.btn_output,
            self.btn_process,
            self.btn_clear,
            self.btn_open_folder,
            self.btn_help,
            self.btn_info,
            self.btn_manual,
            self.btn_about,
            self.btn_log,
            self.btn_report,
            self.btn_validate,
            self.btn_open_last,
            self.btn_reset_pct,
            self.btn_copy_summary,
        ):
            btn.configure(state=state)

        self.metric_estado.config(text="PROCESANDO" if processing else "LISTO")
        self.root.update_idletasks()

    def _build_summary_text(self):
        if not self.last_result:
            return None

        specialty = self.specialty_var.get().strip().lower()

        if specialty == "cardiologia":
            pct_info = (
                f"Consultas/Cuidados: {self.update_consulta_pct_var.get()}%\n"
                f"Procedimientos: {self.update_proced_pct_var.get()}%"
            )
        else:
            pct_info = f"Porcentaje aplicado: {self.update_pct_var.get()}%"

        return (
            "Resumen del último proceso\n\n"
            f"Archivo generado: {self.last_result.output_path}\n"
            f"Especialidad: {self.specialty_var.get().strip().upper()}\n"
            f"PRE: {self.last_result.pre_rows}\n"
            f"PEDIR FAC: {self.last_result.pedir_rows}\n"
            f"AMARILLO: {self.last_result.amarillo_rows}\n"
            f"ANULADOS: {self.last_result.anulados_rows}\n"
            f"{pct_info}"
        )

    def copy_summary(self):
        summary = self._build_summary_text()
        if not summary:
            messagebox.showinfo(
                "Resumen no disponible",
                "Todavía no existe un proceso finalizado para copiar."
            )
            return

        self.root.clipboard_clear()
        self.root.clipboard_append(summary)
        self.root.update()

        self.log("El resumen del último proceso fue copiado al portapapeles.")
        messagebox.showinfo(
            "Resumen copiado",
            "El resumen del último proceso fue copiado al portapapeles correctamente.\n\nAhora puede pegarlo en un correo, informe o mensaje interno."
        )

    def reset_percentages(self):
        self._load_specialty_percentage()
        self.log("Los porcentajes visibles fueron restablecidos a sus valores base.")

        messagebox.showinfo(
            "Porcentajes restablecidos",
            "Los porcentajes visibles del módulo fueron restablecidos a los valores base configurados para la especialidad seleccionada.\n\nNo se realizaron cambios permanentes en la configuración."
        )

    def validate_file(self):
        input_path = self.input_var.get().strip()
        specialty = self.specialty_var.get().strip().lower()

        if not input_path:
            messagebox.showwarning(
                "Archivo no seleccionado",
                "Seleccione primero un archivo de entrada para realizar la validación."
            )
            return

        if not os.path.exists(input_path):
            messagebox.showwarning(
                "Archivo no encontrado",
                "El archivo seleccionado no existe o no está disponible en la ruta indicada."
            )
            return

        required = self._get_required_sheets_for_specialty(specialty)

        try:
            import openpyxl
            wb = openpyxl.load_workbook(input_path, read_only=True, data_only=True)
            sheetnames = wb.sheetnames

            missing = [s for s in required if s not in sheetnames]

            if missing:
                messagebox.showwarning(
                    "Validación incompleta",
                    (
                        "Se encontraron observaciones en el archivo seleccionado.\n\n"
                        f"Especialidad evaluada: {specialty.upper()}\n"
                        "Hojas requeridas faltantes:\n"
                        + "\n".join(f"- {x}" for x in missing)
                        + "\n\nRevise la estructura del archivo antes de procesarlo."
                    ),
                )
                self.log(f"Validación con observaciones. Hojas faltantes: {', '.join(missing)}")
                return

            messagebox.showinfo(
                "Validación completada",
                (
                    "La estructura general del archivo fue validada correctamente para el módulo de facturados.\n\n"
                    f"Especialidad evaluada: {specialty.upper()}\n"
                    "Resultado de la validación:\n"
                    "- Archivo accesible\n"
                    "- Hojas requeridas detectadas\n"
                    "- Estructura mínima compatible\n\n"
                    "Puede continuar con el procesamiento."
                ),
            )
            self.log("Validación del archivo completada correctamente.")

        except Exception as e:
            messagebox.showerror(
                "No fue posible validar el archivo",
                f"Ocurrió un inconveniente durante la validación del archivo.\n\nDetalle técnico:\n{e}"
            )
            self.log(f"ERROR al validar archivo: {e}")

    def process(self):
        input_path = self.input_var.get().strip()
        output_path = self.output_var.get().strip()
        especialidad = self.specialty_var.get().strip().lower()

        if not input_path:
            messagebox.showwarning(
                "Archivo de entrada requerido",
                "Seleccione un archivo Excel de entrada antes de iniciar el procesamiento."
            )
            return

        if not output_path:
            messagebox.showwarning(
                "Archivo de salida requerido",
                "Defina una ruta de salida válida antes de iniciar el procesamiento."
            )
            return

        try:
            payment_data = self._build_payment_payload()
        except Exception as e:
            messagebox.showwarning("Parámetros inválidos", str(e))
            return

        self.log(f"Inicio de procesamiento para especialidad: {especialidad}")

        if payment_data["mode"] == "cardio":
            self.log(
                "Porcentajes aplicados -> "
                f"Consultas/Cuidados: {int(round(payment_data['payment_cfg']['consulta_pct'] * 100))}% | "
                f"Procedimientos: {int(round(payment_data['payment_cfg']['procedimiento_pct'] * 100))}%"
            )
        else:
            self.log(
                f"Porcentaje aplicado al proceso: {int(round(payment_data['payment_pct'] * 100))}%"
            )

        self.progress_var.set(0)
        self.set_status("Iniciando proceso...", 0)
        self.set_processing_state(True)

        try:
            self.root.update()

            process_kwargs = dict(
                input_path=input_path,
                output_path=output_path,
                especialidad=especialidad,
                progress_cb=self.update_progress,
            )

            if payment_data["mode"] == "cardio":
                process_kwargs["payment_cfg"] = payment_data["payment_cfg"]
            else:
                process_kwargs["payment_pct"] = payment_data["payment_pct"]

            result = process_excel(**process_kwargs)
            self.last_result = result
            current_settings = get_specialty_settings(especialidad)

            if payment_data["mode"] == "cardio":
                current_consulta = current_settings.get("consulta_pct", 0.90)
                current_proced = current_settings.get("procedimiento_pct", 0.36)

                new_consulta = payment_data["payment_cfg"]["consulta_pct"]
                new_proced = payment_data["payment_cfg"]["procedimiento_pct"]

                changed = (
                    abs(new_consulta - current_consulta) > 1e-9
                    or abs(new_proced - current_proced) > 1e-9
                )

                if changed:
                    keep_as_base = messagebox.askyesno(
                        "Actualizar porcentajes base",
                        (
                            "Se detectó una diferencia entre los porcentajes base de cardiología y los utilizados en este proceso.\n\n"
                            f"Valores base actuales:\n"
                            f"- Consultas/Cuidados: {int(round(current_consulta * 100))}%\n"
                            f"- Procedimientos: {int(round(current_proced * 100))}%\n\n"
                            f"Valores usados en este proceso:\n"
                            f"- Consultas/Cuidados: {int(round(new_consulta * 100))}%\n"
                            f"- Procedimientos: {int(round(new_proced * 100))}%\n\n"
                            "¿Desea conservar estos nuevos porcentajes como base para futuros procesos?"
                        ),
                    )

                    if keep_as_base:
                        set_specialty_settings(
                            especialidad,
                            {
                                "consulta_pct": new_consulta,
                                "procedimiento_pct": new_proced,
                            },
                        )
                        self._load_specialty_percentage()
                        self.log(
                            "Se actualizaron los porcentajes base de cardiología a "
                            f"{int(round(new_consulta * 100))}% y {int(round(new_proced * 100))}%."
                        )
                    else:
                        self._load_specialty_percentage()
                        self.log("Los porcentajes nuevos se aplicaron únicamente a este proceso.")

            else:
                current_base = current_settings.get("payment_pct", 0.70)
                new_payment = payment_data["payment_pct"]

                if abs(new_payment - current_base) > 1e-9:
                    keep_as_base = messagebox.askyesno(
                        "Actualizar porcentaje base",
                        (
                            "Se detectó una diferencia entre el porcentaje base actual y el utilizado en este proceso.\n\n"
                            f"Porcentaje base actual para {especialidad.upper()}: {int(round(current_base * 100))}%\n"
                            f"Porcentaje usado en este proceso: {int(round(new_payment * 100))}%\n\n"
                            "¿Desea conservar este nuevo porcentaje como base para futuros procesos?"
                        ),
                    )

                    if keep_as_base:
                        set_specialty_settings(
                            especialidad,
                            {
                                "payment_pct": new_payment,
                            },
                        )
                        self._load_specialty_percentage()
                        self.log(
                            f"Se actualizó el porcentaje base de {especialidad} a {int(round(new_payment * 100))}%."
                        )
                    else:
                        self._load_specialty_percentage()
                        self.log("El cambio de porcentaje se aplicó únicamente a este proceso.")

            self.progress_var.set(100)
            self.status_var.set("Proceso finalizado correctamente.")
            self.log("Proceso finalizado con éxito.")
            self.log(f"Archivo generado: {result.output_path}")
            self.log(
                f"Resumen -> PRE: {result.pre_rows} | PEDIR FAC: {result.pedir_rows} | "
                f"AMARILLO: {result.amarillo_rows} | ANULADOS: {result.anulados_rows}"
            )

            self.metric_pre.config(text=str(result.pre_rows))
            self.metric_pedir.config(text=str(result.pedir_rows))
            self.metric_estado.config(text="FINALIZADO")

            if payment_data["mode"] == "cardio":
                pct_info = (
                    f"Consultas/Cuidados: {int(round(payment_data['payment_cfg']['consulta_pct'] * 100))}%\n"
                    f"Procedimientos: {int(round(payment_data['payment_cfg']['procedimiento_pct'] * 100))}%"
                )
            else:
                pct_info = f"Porcentaje aplicado: {int(round(payment_data['payment_pct'] * 100))}%"

            messagebox.showinfo(
                "Procesamiento completado",
                (
                    "El procesamiento finalizó correctamente.\n\n"
                    "Resultado general del proceso:\n"
                    f"- Archivo generado: {result.output_path}\n"
                    f"- Registros en PRE: {result.pre_rows}\n"
                    f"- Registros en PEDIR FAC: {result.pedir_rows}\n"
                    f"- Registros en AMARILLO: {result.amarillo_rows}\n"
                    f"- Registros en ANULADOS: {result.anulados_rows}\n"
                    f"- Parámetros usados:\n{pct_info}\n\n"
                    "Puede revisar el resultado desde “Abrir ubicación del archivo” o desde “Abrir último generado”."
                ),
            )

        except Exception as e:
            self.status_var.set("Ocurrió un error durante el proceso.")
            self.metric_estado.config(text="ERROR")
            self.log(f"ERROR: {str(e)}")
            messagebox.showerror(
                "Procesamiento interrumpido",
                f"Ocurrió un inconveniente durante la ejecución del módulo.\n\nDetalle técnico:\n{e}"
            )

        finally:
            self.set_processing_state(False)

    def show_help(self):
        messagebox.showinfo(
            "Ayuda de uso del módulo de facturados",
            (
                "Este módulo permite procesar archivos de honorarios médicos facturados por especialidad.\n\n"
                "Pasos recomendados:\n"
                "1. Seleccione la especialidad correspondiente.\n"
                "2. Revise los porcentajes base mostrados en pantalla.\n"
                "3. Ajuste los porcentajes solo si el caso lo requiere.\n"
                "4. Seleccione el archivo Excel de entrada.\n"
                "5. Verifique o modifique la ruta de salida.\n"
                "6. Use “Validar archivo” si desea revisar la estructura antes del proceso.\n"
                "7. Pulse “Procesar archivo”.\n"
                "8. Revise la bitácora y el archivo generado."
            ),
        )

    def show_info(self):
        messagebox.showinfo(
            "Guía del módulo de facturados",
            (
                "El módulo de facturados automatiza el análisis y cálculo de honorarios médicos a partir de archivos institucionales en Excel.\n\n"
                "Especialidades implementadas actualmente:\n"
                "- Ortopedia\n"
                "- Urología\n"
                "- Cardiología\n"
                "- Fonoaudiología\n\n"
                "Funciones principales del módulo:\n"
                "- Lectura de archivos Excel\n"
                "- Validación básica de hojas requeridas\n"
                "- Aplicación de reglas de exclusión\n"
                "- Detección de anulados\n"
                "- Detección de duplicados, cuando aplique\n"
                "- Cálculo automático de valores\n"
                "- Generación de hojas PRE, PEDIR FAC, AMARILLO, ANULADOS y otras según la especialidad"
            ),
        )

    def show_about(self):
        messagebox.showinfo(
            "Acerca del sistema",
            (
                f"{APP_TITLE}\n\n"
                "Versión de desarrollo\n"
                "Sistema de automatización para el procesamiento de honorarios médicos por especialidad.\n\n"
                "Módulo activo:\n"
                "Facturados\n\n"
                "Especialidades implementadas actualmente:\n"
                "- Ortopedia\n"
                "- Urología\n"
                "- Cardiología\n"
                "- Fonoaudiología\n\n"
                "Tecnologías principales:\n"
                "- Python\n"
                "- Tkinter\n"
                "- OpenPyXL"
            ),
        )

    def focus_log(self):
        self.log_text.focus_set()
        self.log_text.see("end")
        self.log("Se abrió la bitácora del sistema para revisión.")

    def quick_report(self):
        summary = self._build_summary_text()

        if not summary:
            messagebox.showinfo(
                "Resumen no disponible",
                "Todavía no existe un procesamiento finalizado para mostrar."
            )
            return

        messagebox.showinfo("Resumen del último proceso", summary)
        self.log("Se consultó el resumen del último proceso.")


def run_app():
    root = tk.Tk()
    MainWindow(root)
    root.mainloop()
