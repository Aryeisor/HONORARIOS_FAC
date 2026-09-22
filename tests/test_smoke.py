from datetime import datetime

from openpyxl import Workbook, load_workbook

from honorarios_app.core.processor import process_excel
from honorarios_app.specialties.cardiologia import CardiologiaRules
from honorarios_app.specialties.fonoaudiologia import FonoaudiologiaRules
from honorarios_app.specialties.ortopedia import OrtopediaRules
from honorarios_app.specialties.urologia import UrologiaRules
from honorarios_app.services.specialty_factory import get_specialty


def _base_headers():
    return [
        "NO FAC",
        "IDENTIFICACION PACIENTE",
        "FECHA DEL SERVICIO",
        "PROCEDIMIENTO",
        "SALDO",
    ]


def _anulaciones_sheet(rows):
    wb = Workbook()
    ws = wb.active
    headers = [
        "NO FAC",
        "IDENTIFICACION PACIENTE",
        "PROCEDIMIENTO",
        "SALDO",
    ]
    ws.append(headers)
    for row in rows:
        ws.append(row)
    col_map = {header: column for column, header in enumerate(headers, start=1)}
    return ws, col_map


def test_factory_replaces_gastroenterologia_with_fonoaudiologia():
    assert isinstance(get_specialty("fonoaudiologia"), FonoaudiologiaRules)
    try:
        get_specialty("gastroenterologia")
    except ValueError:
        pass
    else:
        raise AssertionError("Gastroenterologia no debe seguir registrada")


def test_cardiologia_detects_anulacion_with_different_invoices():
    ws, col_map = _anulaciones_sheet([
        ["173322", "901351767", "Electrocardiograma", -90326],
        ["174615", "901351767", "Electrocardiograma", 90326],
    ])

    anulable_rows = CardiologiaRules().detect_anulable_rows(ws, 1, col_map)

    assert anulable_rows == {2, 3}


def test_fonoaudiologia_reuses_cardiologia_anulados_logic():
    ws, col_map = _anulaciones_sheet([
        ["A1", "123", "Terapia", 100],
        ["A2", "123", "Terapia", -100],
    ])

    assert FonoaudiologiaRules().detect_anulable_rows(ws, 1, col_map) == {2, 3}


def test_fonoaudiologia_duplicate_rules():
    rows = [
        ["I1", "123", "2026-06-01 08:00", "INTERCONSULTA", 120, 0],
        ["T1", "123", "2026-06-01 09:00", "TERAPIA FONOAUDIOLOGICA", 80, 0],
        ["T2", "123", "2026-06-01 10:00", "TERAPIA FONOAUDIOLOGICA", 100, 0],
        ["I2", "456", "2026-06-02", "INTERCONSULTA", 50, 0],
        ["I3", "456", "2026-06-02", "INTERCONSULTA", 70, 0],
    ]
    final_pre, final_pedir, duplicates = (
        FonoaudiologiaRules().filter_duplicate_consult_rows(
            rows, [], [row[:5] for row in rows], [], _base_headers()
        )
    )

    assert [row[0] for row in final_pre] == ["T2", "I3"]
    assert [row[0] for row in final_pedir] == ["T2", "I3"]
    assert [row[0] for row in duplicates] == ["I1", "T1", "I2"]


def test_fonoaudiologia_calculation_uses_937001_and_preserves_interconsult():
    rules = FonoaudiologiaRules()
    common = dict(
        row_values=[], contrato_value="COOSALUD00125", no_fac_value="F1",
        saldo_value=1000, tarifario={"937001": {"cir": 1000, "ayd2": 0}},
        hosvi_by_fact={}, hosvi_by_alt={}, ws_original=None, row_number=2,
        col_map={}, idx_vlr_aut=1, hosvi_cursor_fact={}, hosvi_cursor_alt={},
        payment_pct=0.65,
    )
    therapy = rules.process_row(
        procedimiento_value="TERAPIA FONOAUDIOLOGICA INTEGRAL",
        cod_proc_value="937001", **common
    )
    interconsult = rules.process_row(
        procedimiento_value="INTERCONSULTA", cod_proc_value="890402", **common
    )

    assert therapy["calc_values"] == [1000, 650, 650, 350]
    assert therapy["nuevo_vlr_aut"] == 650
    assert interconsult["calc_values"] == [0, 0, 0, 0]
    assert interconsult["nuevo_vlr_aut"] is None


def test_fonoaudiologia_integrates_without_hosvireport(tmp_path):
    input_path = tmp_path / "entrada_fono.xlsx"
    output_path = tmp_path / "salida_fono.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "FACTURADOS"
    ws.append([
        "CONTRATO", "NO FAC", "CODIGO PROCEDIMIENTO", "PROCEDIMIENTO",
        "SALDO", "IDENTIFICACION PACIENTE", "NOMBRE PACIENTE",
        "FECHA DEL SERVICIO", "VLR. A AUTORIZAR", "TIPO FAC",
    ])
    ws.append(["COOSALUD00125", "I1", "890402", "INTERCONSULTA", 900,
               "123", "PACIENTE", "2026-06-01", 777, ""])
    ws.append(["COOSALUD00125", "T1", "937001",
               "TERAPIA FONOAUDIOLÓGICA INTEGRAL", 800, "123", "PACIENTE",
               "2026-06-01", 0, ""])
    ws.append(["COOSALUD00125", "T2", "937001",
               "TERAPIA FONOAUDIOLÓGICA INTEGRAL", 1000, "123", "PACIENTE",
               "2026-06-01", 0, ""])
    coosalud = wb.create_sheet("COOSALUD")
    coosalud.append(["CUPS", "CIRUJANO", "AYUDANTE 2"])
    coosalud.append(["937001", 1000, 0])
    wb.save(input_path)

    process_excel(
        str(input_path), str(output_path), especialidad="fonoaudiologia",
        payment_pct=0.65,
    )
    result = load_workbook(output_path, data_only=False)

    assert result.sheetnames == ["PRE", "AMARILLO", "PEDIR FAC", "ANULADOS", "DUPLICADOS"]
    assert result["PRE"].cell(1, 11).value == "Valor COOSALUD"
    assert result["PRE"].cell(1, 12).value == "65% COOS"
    assert result["PRE"].cell(2, 2).value == "T2"
    assert result["PRE"].cell(2, 9).value == 650
    assert result["PEDIR FAC"].cell(2, 2).value == "T2"
    assert {result["DUPLICADOS"].cell(row, 2).value for row in (2, 3)} == {"I1", "T1"}


def test_cardiologia_keeps_primary_anulacion_priority():
    ws, col_map = _anulaciones_sheet([
        ["A1", "123", "Electrocardiograma", 100],
        ["A1", "123", "Electrocardiograma", -100],
        ["B1", "123", "Electrocardiograma", 50],
    ])

    anulable_rows = CardiologiaRules().detect_anulable_rows(ws, 1, col_map)

    assert anulable_rows == {2, 3}


def test_ortopedia_detects_anulacion_with_different_invoices():
    ws, col_map = _anulaciones_sheet([
        ["173322", "901351767", "Electrocardiograma", -90326],
        ["174615", "901351767", "Electrocardiograma", 90326],
    ])

    anulable_rows = OrtopediaRules().detect_anulable_rows(ws, 1, col_map)

    assert anulable_rows == {2, 3}


def test_urologia_detects_anulacion_with_different_invoices_and_code():
    wb = Workbook()
    ws = wb.active
    headers = [
        "NO FAC",
        "IDENTIFICACION PACIENTE",
        "CODIGO PROCEDIMIENTO",
        "PROCEDIMIENTO",
        "SALDO",
        "FECHA DEL SERVICIO",
    ]
    ws.append(headers)
    ws.append(["167116", "13458021", "890402", "INTERCONSULTA", -55610, "2026-06-01"])
    ws.append(["167533", "13458021", "890402", "INTERCONSULTA", 55610, "2026-06-20"])
    col_map = {header: column for column, header in enumerate(headers, start=1)}

    rules = UrologiaRules()

    assert rules.detect_anulable_rows(ws, 1, col_map) == {2, 3}
    assert rules.detect_duplicate_rows(ws, 1, col_map) == {}


def test_ortopedia_filters_pre_and_pedir_fact_with_same_duplicate_indexes():
    pre_rows = [
        ["F1", "12345", "2026-05-08 08:00:00", "Consulta", 95000, 1],
        ["F2", "12345", "2026-05-08 14:00:00", "Interconsulta", 70000, 2],
        ["F3", "12345", "2026-05-08 16:00:00", "Cuidado", 82000, 3],
    ]
    pedir_rows = [row[:5] for row in pre_rows]

    final_pre, final_pedir, duplicate_rows = OrtopediaRules().filter_duplicate_consult_rows(
        pre_rows,
        [],
        pedir_rows,
        [],
        _base_headers(),
    )

    assert [row[0] for row in final_pre] == ["F1"]
    assert [row[0] for row in final_pedir] == ["F1"]
    assert [row[0] for row in duplicate_rows] == ["F2", "F3"]


def test_cardiologia_filter_keeps_highest_saldo_by_patient_and_date():
    pre_rows_coo = [
        ["F1", "12345", datetime(2026, 5, 8, 8), "Consulta", 95000, 1],
        ["F2", "12345", datetime(2026, 5, 8, 14), "Interconsulta", 70000, 2],
    ]
    pre_rows_other = [
        ["F3", "12345", datetime(2026, 5, 8, 16), "Cuidado", 82000, 3],
    ]

    final_rows, final_pedir_rows, duplicate_rows = CardiologiaRules().filter_duplicate_consult_rows(
        pre_rows_coo,
        pre_rows_other,
        [row[:5] for row in pre_rows_coo],
        [row[:5] for row in pre_rows_other],
        _base_headers(),
    )

    assert [row[0] for row in final_rows] == ["F1"]
    assert [row[0] for row in final_pedir_rows] == ["F1"]
    assert [row[0] for row in duplicate_rows] == ["F2", "F3"]
    assert duplicate_rows[0][-1] == "CONSULTA DUPLICADA MISMO PACIENTE Y FECHA"
    assert len(duplicate_rows[0]) == len(_base_headers()) + 1


def test_cardiologia_filter_ignores_other_procedures_and_different_dates():
    pre_rows_coo = [
        ["F1", "12345", "2026-05-08 08:00:00", "Consulta", 90000, 1],
        ["F2", "12345", "2026-05-08 10:00:00", "Ecocardiograma", 100000, 2],
        ["F3", "12345", "2026-05-09 08:00:00", "Interconsulta", 60000, 3],
    ]

    final_rows, final_pedir_rows, duplicate_rows = CardiologiaRules().filter_duplicate_consult_rows(
        pre_rows_coo,
        [],
        [row[:5] for row in pre_rows_coo],
        [],
        _base_headers(),
    )

    assert len(final_rows) == 3
    assert len(final_pedir_rows) == 3
    assert duplicate_rows == []


def test_cardiologia_filter_does_not_mutate_input_lists():
    pre_rows_coo = [
        ["F1", "12345", "2026-05-08", "Consulta", 90000, 1],
        ["F2", "12345", "2026-05-08", "Interconsulta", 60000, 2],
    ]
    original_rows = [list(row) for row in pre_rows_coo]

    CardiologiaRules().filter_duplicate_consult_rows(
        pre_rows_coo,
        [],
        [row[:5] for row in pre_rows_coo],
        [],
        _base_headers(),
    )

    assert pre_rows_coo == original_rows


def test_cardiologia_anulados_have_priority_over_duplicates(tmp_path):
    input_path = tmp_path / "entrada.xlsx"
    output_path = tmp_path / "salida.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "FACTURADOS"
    ws.append([
        "CONTRATO",
        "NO FAC",
        "CODIGO PROCEDIMIENTO",
        "PROCEDIMIENTO",
        "SALDO",
        "IDENTIFICACION PACIENTE",
        "NOMBRE PACIENTE",
        "FECHA DEL SERVICIO",
        "ESTADO",
        "VLR. A AUTORIZAR",
        "TIPO FAC",
    ])
    ws.append(["COOSALUD00225", "A1", "C1", "Consulta", 100, "123", "PACIENTE", "2026-05-08", "", 0, ""])
    ws.append(["COOSALUD00225", "A1", "C1", "Consulta", -100, "123", "PACIENTE", "2026-05-08", "", 0, ""])
    ws.append(["COOSALUD00225", "F1", "C1", "Consulta", 90000, "123", "PACIENTE", "2026-05-08", "", 0, ""])
    ws.append(["COOSALUD00225", "F2", "C1", "Interconsulta", 60000, "123", "PACIENTE", "2026-05-08", "", 0, ""])

    ws_coosalud = wb.create_sheet("COOSALUD")
    ws_coosalud.append(["CUPS", "CIRUJANO", "AYUDANTE 2"])
    ws_coosalud.append(["C1", 100000, 0])
    wb.save(input_path)

    process_excel(str(input_path), str(output_path), especialidad="cardiologia")

    result = load_workbook(output_path, data_only=False)
    assert result["ANULADOS"].max_row == 3
    assert result["PRE"].cell(2, 2).value == "F1"
    assert result["PEDIR FAC"].cell(2, 2).value == "F1"
    assert result["PEDIR FAC"].max_row == 4
    assert result["DUPLICADOS"].max_row == 2
    assert result["DUPLICADOS"].cell(2, 2).value == "F2"


def test_ortopedia_integrates_anulados_duplicates_pre_and_pedir(tmp_path):
    input_path = tmp_path / "entrada_ortopedia.xlsx"
    output_path = tmp_path / "salida_ortopedia.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "FACTURADOS"
    ws.append([
        "CONTRATO",
        "NO FAC",
        "CODIGO PROCEDIMIENTO",
        "PROCEDIMIENTO",
        "SALDO",
        "IDENTIFICACION PACIENTE",
        "NOMBRE PACIENTE",
        "FECHA DEL SERVICIO",
        "ESTADO",
        "VLR. A AUTORIZAR",
        "TIPO FAC",
    ])
    ws.append(["COOSALUD00225", "A1", "E1", "Electrocardiograma", -100, "123", "PACIENTE", "2026-05-08", "", 0, ""])
    ws.append(["COOSALUD00225", "A2", "E1", "Electrocardiograma", 100, "123", "PACIENTE", "2026-05-08", "", 0, ""])
    ws.append(["COOSALUD00225", "F1", "C1", "Consulta", 90000, "123", "PACIENTE", "2026-05-08", "", 0, ""])
    ws.append(["COOSALUD00225", "F2", "C1", "Interconsulta", 60000, "123", "PACIENTE", "2026-05-08", "", 0, ""])

    ws_hosvi = wb.create_sheet("HOSVIREPORT")
    ws_hosvi.append(["FACTURA", "COD_PROCED"])
    ws_coosalud = wb.create_sheet("COOSALUD")
    ws_coosalud.append(["CUPS", "CIRUJANO", "AYUDANTE 2"])
    ws_coosalud.append(["C1", 100000, 0])
    ws_coosalud.append(["E1", 100000, 0])
    wb.save(input_path)

    process_excel(str(input_path), str(output_path), especialidad="ortopedia")

    result = load_workbook(output_path, data_only=False)
    assert result["ANULADOS"].max_row == 3
    assert result["PRE"].cell(2, 2).value == "F1"
    assert result["PEDIR FAC"].cell(2, 2).value == "F1"
    assert result["PEDIR FAC"].max_row == 4
    assert result["DUPLICADOS"].max_row == 2
    assert result["DUPLICADOS"].cell(2, 2).value == "F2"


def test_urologia_anulados_have_priority_over_duplicates_with_different_invoices(tmp_path):
    input_path = tmp_path / "entrada_urologia.xlsx"
    output_path = tmp_path / "salida_urologia.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "FACTURADOS"
    ws.append([
        "CONTRATO",
        "NO FAC",
        "CODIGO PROCEDIMIENTO",
        "PROCEDIMIENTO",
        "SALDO",
        "IDENTIFICACION PACIENTE",
        "NOMBRE PACIENTE",
        "FECHA DEL SERVICIO",
        "ESTADO",
        "VLR. A AUTORIZAR",
        "TIPO FAC",
    ])
    ws.append(["COOSALUD00225", "167116", "890402", "INTERCONSULTA", -55610, "13458021", "PACIENTE", "2026-06-01", "", 0, ""])
    ws.append(["COOSALUD00225", "167533", "890402", "INTERCONSULTA", 55610, "13458021", "PACIENTE", "2026-06-20", "", 0, ""])

    ws_hosvi = wb.create_sheet("HOSVIREPORT")
    ws_hosvi.append(["FACTURA", "COD_PROCED"])
    ws_coosalud = wb.create_sheet("COOSALUD")
    ws_coosalud.append(["CUPS", "CIRUJANO", "AYUDANTE 2"])
    ws_coosalud.append(["890402", 100000, 0])
    wb.save(input_path)

    process_excel(str(input_path), str(output_path), especialidad="urologia")

    result = load_workbook(output_path, data_only=False)
    assert result["ANULADOS"].max_row == 3
    assert {result["ANULADOS"].cell(row, 2).value for row in (2, 3)} == {"167116", "167533"}
    assert result["DUPLICADOS"].max_row == 1
    assert result["PRE"].max_row == 1
    assert result["PEDIR FAC"].max_row == 1
